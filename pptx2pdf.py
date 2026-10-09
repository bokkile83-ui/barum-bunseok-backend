# -*- coding: utf-8 -*-
"""★v823 제223조 — 보장분석지 PPT(값 주입 완료본) → A3 PDF.
Railway 에 LibreOffice 가 없어 PPTX→PDF 변환이 불가하다. 폼은 사각형·타원·텍스트상자뿐이므로
도형 좌표·채움·선·글자(색·크기·굵기·정렬)를 그대로 HTML 절대좌표로 옮겨 weasyprint 로 PDF 를 만든다.
PDF 값 = PPT 값(같은 파일을 읽는다 — 등식 유지). 글꼴: 서버에 HY엽서M 이 없어 NanumGothic/NotoSansKR 대체."""
import html as _h

from pptx import Presentation
from pptx.util import Emu
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

_CM = 360000.0


def _rgb_of(fill):
    try:
        if fill.type is not None and 'SOLID' in str(fill.type):  # MSO_FILL.SOLID
            return '#' + str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def _grad_of(shape):
    """선형 그라데이션(머리선) — 첫/끝 색만."""
    try:
        gf = shape._element.spPr.find(qn('a:gradFill'))
        if gf is None: return None
        cols = [c.find(qn('a:srgbClr')).get('val') for c in gf.find(qn('a:gsLst')).findall(qn('a:gs'))]
        if len(cols) >= 2: return 'linear-gradient(90deg,#%s,#%s)' % (cols[0], cols[-1])
    except Exception:
        return None
    return None


def _line_of(shape):
    try:
        ln = shape._element.spPr.find(qn('a:ln'))
        if ln is None: return None
        if ln.find(qn('a:noFill')) is not None: return None
        sf = ln.find(qn('a:solidFill'))
        if sf is None: return None
        c = sf.find(qn('a:srgbClr'))
        if c is None: return None
        w = int(ln.get('w') or 9525) / 12700.0  # pt
        return (w, '#' + c.get('val'))
    except Exception:
        return None


def _run_html(r, base_pt):
    t = _h.escape(r.text or '')
    if not t: return ''
    sz = r.font.size.pt if r.font.size else base_pt
    col = None
    try:
        if r.font.color and r.font.color.type is not None and r.font.color.rgb is not None:
            col = '#' + str(r.font.color.rgb)
    except Exception:
        col = None
    st = 'font-size:%.1fpt;' % sz
    if col: st += 'color:%s;' % col
    if r.font.bold: st += 'font-weight:700;'
    # 노랑 하이라이트(CI)
    try:
        hl = r._r.find(qn('a:rPr'))
        if hl is not None and hl.find(qn('a:highlight')) is not None: st += 'background:#FFFF00;'
    except Exception:
        pass
    t = t.replace('\n', '<br>')
    return '<span style="%s">%s</span>' % (st, t)


def _tf_html(shape):
    tf = shape.text_frame
    bp = tf._txBody.find(qn('a:bodyPr'))
    ins = {k: int(bp.get(k)) / _CM if bp is not None and bp.get(k) else 0.25 for k in ('lIns', 'rIns', 'tIns', 'bIns')}
    for k, d in (('lIns', 0.25), ('rIns', 0.25), ('tIns', 0.127), ('bIns', 0.127)):
        if bp is None or not bp.get(k): ins[k] = d
    anchor = (bp.get('anchor') if bp is not None else None) or 't'
    wrap = (bp.get('wrap') if bp is not None else None) != 'none'
    paras = []
    for p in tf.paragraphs:
        al = {1: 'left', 2: 'center', 3: 'right'}.get(p.alignment, 'left') if p.alignment is not None else 'left'
        lh = ''
        try:
            if p.line_spacing is not None and hasattr(p.line_spacing, 'pt'):
                lh = 'line-height:%.1fpt;' % p.line_spacing.pt
            elif isinstance(p.line_spacing, float):
                lh = 'line-height:%.2f;' % p.line_spacing
        except Exception:
            lh = ''
        inner = ''.join(_run_html(r, 10) for r in p.runs) or '&nbsp;'
        paras.append('<div style="text-align:%s;%s">%s</div>' % (al, lh, inner))
    # ★v824 (지점장 2026.10.09 「칸들에 비해 글자가 내려간다」): weasyprint 의 flex 세로 가운데는 글자가 칸보다 클 때
    #   위로 붙어 버린다(음수 여유를 0으로 본다). → 절대좌표 + translateY 로 가운데/아래를 잡는다. 줄간격 1.15 고정.
    if anchor == 'ctr':
        pos = 'top:50%%;transform:translateY(-50%%);padding:0 %.2fcm 0 %.2fcm;' % (ins['rIns'], ins['lIns'])
    elif anchor == 'b':
        pos = 'bottom:%.2fcm;padding:0 %.2fcm 0 %.2fcm;' % (ins['bIns'], ins['rIns'], ins['lIns'])
    else:
        pos = 'top:%.2fcm;padding:0 %.2fcm 0 %.2fcm;' % (ins['tIns'], ins['rIns'], ins['lIns'])
    return ('<div style="position:absolute;left:0;right:0;line-height:1.15;%s%s">%s</div>'
            % (pos, '' if wrap else 'white-space:nowrap;', ''.join(paras)))


def pptx_to_html(src, font_family="'HY엽서M','NanumGothic','Noto Sans KR','Malgun Gothic',sans-serif"):
    prs = Presentation(src)
    W, H = prs.slide_width / _CM, prs.slide_height / _CM
    sl = prs.slides[0]
    out = []
    for sh in sl.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            # ★v825 폼 v8: 디자인 그림이 배경이다 — 그대로 PDF에 싣는다
            try:
                import base64 as _b64
                _img = sh.image; _mime = _img.content_type or 'image/png'
                x, y, w, h = sh.left / _CM, sh.top / _CM, sh.width / _CM, sh.height / _CM
                out.append('<img style="position:absolute;left:%.3fcm;top:%.3fcm;width:%.3fcm;height:%.3fcm" src="data:%s;base64,%s">'
                           % (x, y, w, h, _mime, _b64.b64encode(_img.blob).decode()))
            except Exception:
                pass
            continue
        x, y, w, h = sh.left / _CM, sh.top / _CM, sh.width / _CM, sh.height / _CM
        st = 'position:absolute;left:%.3fcm;top:%.3fcm;width:%.3fcm;height:%.3fcm;box-sizing:border-box;' % (x, y, w, h)
        is_oval = False
        try:
            is_oval = (sh.auto_shape_type is not None and 'OVAL' in str(sh.auto_shape_type))
        except Exception:
            is_oval = False
        if is_oval: st += 'border-radius:50%;'
        if sh.shape_type in (MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX):
            g = _grad_of(sh)
            if g: st += 'background:%s;' % g
            else:
                c = _rgb_of(sh.fill)
                if c: st += 'background:%s;' % c
            ln = _line_of(sh)
            if ln: st += 'border:%.2fpt solid %s;' % (max(ln[0], 0.4), ln[1])
        elif sh.shape_type == MSO_SHAPE_TYPE.LINE:
            ln = _line_of(sh)
            if ln: st += 'border-top:%.2fpt solid %s;height:0;' % (ln[0], ln[1])
        inner = _tf_html(sh) if sh.has_text_frame and sh.text_frame.text.strip() else ''
        out.append('<div style="%s">%s</div>' % (st, inner))
    html = ('<!doctype html><html><head><meta charset="utf-8"><style>'
            '@page{size:%.2fcm %.2fcm;margin:0}'
            'html,body{margin:0;padding:0}'
            'body{font-family:%s;color:#111;}'
            '.pg{position:relative;width:%.2fcm;height:%.2fcm;overflow:hidden;background:#fff}'
            '</style></head><body><div class="pg">%s</div></body></html>'
            % (W, H, font_family, W, H, ''.join(out)))
    return html


def pptx_to_pdf(src, out_pdf):
    from weasyprint import HTML
    html = pptx_to_html(src)
    HTML(string=html).write_pdf(out_pdf)
    return out_pdf


if __name__ == '__main__':
    import sys
    pptx_to_pdf(sys.argv[1], sys.argv[2]); print('ok', sys.argv[2])
