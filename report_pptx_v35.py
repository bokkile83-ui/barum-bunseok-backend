# -*- coding: utf-8 -*-
"""BARUM 보장진단서 PPT v35 — 뼈대 그대로, 결과값 칸만 입력·수정 가능.
   지점장 지시(2026.07.08): 07.07 디자인을 100% 유지하고, 값이 들어가는 칸에만 입력이 되게 하라.

   방식:
     ① report_weasy 로 보장설명서 PDF 생성            ← 디자인 원본
     ② 각 페이지를 고해상 이미지로 렌더 → 슬라이드 배경  ← 뼈대·라벨·표·게이지 그대로
     ③ rep 에서 '결과값' 문자열 목록을 만든다
     ④ pdftotext -bbox-layout 으로 그 값 단어의 좌표만 찾는다
     ⑤ 그 자리만 배경색으로 지우고 텍스트 상자를 얹는다

   → 라벨·제목·표·심혈관 분류는 손대지 않는다. 값만 클릭·수정된다.
   추가 패키지 없음 (poppler = pdftotext·pdftoppm, 이미 Dockerfile에 있음).
   pdftotext 실패 시 자동으로 종전 이미지 방식으로 폴백 → 앱 안 죽음.
"""
import os, subprocess, tempfile
import xml.etree.ElementTree as ET
from collections import Counter

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

NS = {'x': 'http://www.w3.org/1999/xhtml'}
EMU_PER_PT = 12700
DPI = 200
FONT = "맑은 고딕"


def _setfont(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", name)


def _valueset(rep):
    """rep 에서 '결과값' 문자열만 모은다. 공백 제거 후 비교."""
    V = set()

    def add(x):
        if x is None:
            return
        s = str(x).strip().replace(' ', '')
        if s and s not in ('-', '0'):
            V.add(s)

    add(f"{rep.get('n_contract', 0)}건")
    add(f"{rep.get('premium', 0):,}원")
    add(f"{rep.get('renew', 0)}/{rep.get('nonrenew', 0)}")
    add(f"{rep.get('gap_count', 0)}영역")
    for c in rep.get('coverage', []):
        for i in c.get('items', []):
            add(i.get('v'))
    ci = rep.get('ci', {}) or {}
    for i in ci.get('items', []):
        add(i.get('v'))
    add(ci.get('samang')); add(ci.get('residual'))
    if ci.get('rate'):
        add(f"{ci['rate']}%"); add(f"{100 - ci['rate']}%")
    for c in rep.get('chiryo', []):
        add(c.get('value'))
    for d in rep.get('donuts', []):
        add(f"{d.get('pct')}%")
    for d in rep.get('donut_detail', []):
        add(d.get('have')); add(d.get('rec')); add(f"{d.get('pct')}%")
    for c in rep.get('renew_list', []) + rep.get('nonrenew_list', []):
        add(c.get('v'))
    for i in rep.get('p5_own', []):
        add(i.get('v'))
    return V


def _value_boxes(xml_path, V):
    """페이지별 [(값, xMin, yMin, xMax, yMax)] — 값에 해당하는 단어(묶음)만"""
    root = ET.parse(xml_path).getroot()
    out = []
    for pg in root.findall('.//x:page', NS):
        pw, ph = float(pg.get('width')), float(pg.get('height'))
        found = []
        for ln in pg.findall('.//x:line', NS):
            ws = [(w.text or '', float(w.get('xMin')), float(w.get('yMin')),
                   float(w.get('xMax')), float(w.get('yMax')))
                  for w in ln.findall('x:word', NS)]
            n = len(ws); i = 0
            while i < n:
                for L in (4, 3, 2, 1):                      # 긴 묶음 우선 ('1억2,310만' > '1억')
                    if i + L > n:
                        continue
                    grp = ws[i:i + L]
                    if ''.join(g[0] for g in grp) in V:
                        txt = ' '.join(g[0] for g in grp)
                        found.append((txt, grp[0][1], min(g[2] for g in grp),
                                      grp[-1][3], max(g[4] for g in grp)))
                        i += L
                        break
                else:
                    i += 1
        out.append((pw, ph, found))
    return out


def _fg_of(img, bx):
    """글자색 = bbox 안에서 주변 배경색과 가장 먼 픽셀 (흰 글자·컬러 글자 대응)"""
    x0, y0, x1, y1 = bx
    W, H = img.size
    x0, y0 = max(0, x0), max(0, y0); x1, y1 = min(W, x1), min(H, y1)
    if x1 <= x0 or y1 <= y0:
        return (0, 0, 0)
    ring = []
    for xx in range(x0, x1, max(1, (x1 - x0) // 10)):
        for yy in (max(0, y0 - 2), min(H - 1, y1 + 1)):
            ring.append(img.getpixel((xx, yy))[:3])
    bg = Counter(ring).most_common(1)[0][0] if ring else (255, 255, 255)
    px = list(img.crop((x0, y0, x1, y1)).convert('RGB').getdata())
    return max(px, key=lambda c: (c[0]-bg[0])**2 + (c[1]-bg[1])**2 + (c[2]-bg[2])**2)


def _ink(img, bx):
    """잉크 밀도 → 굵기 판정"""
    x0, y0, x1, y1 = bx
    W, H = img.size
    x0, y0 = max(0, x0), max(0, y0); x1, y1 = min(W, x1), min(H, y1)
    if x1 <= x0 or y1 <= y0:
        return 0.0
    px = list(img.crop((x0, y0, x1, y1)).convert('L').getdata())
    if not px:
        return 0.0
    px.sort(); bg = px[len(px) // 2]
    return sum(1 for v in px if abs(v - bg) > 60) / len(px)


def _erase(img, bx):
    """가로 한 줄씩 좌우 바깥 픽셀 색으로 채운다 → 색 띠·배지·표 배경 보존"""
    x0, y0, x1, y1 = bx
    W, H = img.size
    x0, y0 = max(0, x0), max(0, y0); x1, y1 = min(W, x1), min(H, y1)
    if x1 <= x0 or y1 <= y0:
        return
    lx, rx = max(0, x0 - 2), min(W - 1, x1 + 1)
    n = x1 - x0
    for yy in range(y0, y1):
        lc = img.getpixel((lx, yy))[:3]
        rc = img.getpixel((rx, yy))[:3]
        if lc == rc:
            for xx in range(x0, x1):
                img.putpixel((xx, yy), lc)
        else:
            for i, xx in enumerate(range(x0, x1)):
                t = i / max(1, n - 1)
                img.putpixel((xx, yy), tuple(int(lc[k] + (rc[k] - lc[k]) * t) for k in range(3)))


def build_report_pptx(rep, out, dpi=DPI):
    from report_weasy import build_report_pdf
    from pdf2image import convert_from_path

    tmp = tempfile.mkdtemp()
    pdf = os.path.join(tmp, 'rep.pdf')
    build_report_pdf(rep, pdf)

    try:
        xml = os.path.join(tmp, 'bb.xml')
        subprocess.run(['pdftotext', '-bbox-layout', pdf, xml], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        pages = _value_boxes(xml, _valueset(rep))
    except Exception:
        pages = []                                   # 폴백 = 종전 이미지 방식

    imgs = convert_from_path(pdf, dpi=dpi)

    prs = Presentation()
    pw_pt, ph_pt = (pages[0][0], pages[0][1]) if pages else (595.276, 841.890)
    prs.slide_width = Emu(int(pw_pt * EMU_PER_PT))
    prs.slide_height = Emu(int(ph_pt * EMU_PER_PT))
    BL = prs.slide_layouts[6]

    for idx, im in enumerate(imgs):
        im = im.convert('RGB')
        vals = pages[idx][2] if idx < len(pages) else []
        sx = im.width / pw_pt
        sy = im.height / ph_pt

        meta = []
        for txt, x0, y0, x1, y1 in vals:
            bx = (int(x0 * sx) - 1, int(y0 * sy) - 1, int(x1 * sx) + 2, int(y1 * sy) + 1)
            meta.append((txt, x0, y0, x1, y1, _fg_of(im, bx), _ink(im, bx), bx))
        for *_, bx in meta:
            _erase(im, bx)

        ip = os.path.join(tmp, f'p{idx}.png')
        im.save(ip, 'PNG')

        s = prs.slides.add_slide(BL)
        s.shapes.add_picture(ip, 0, 0, prs.slide_width, prs.slide_height)

        for txt, x0, y0, x1, y1, fg, ink, _bx in meta:
            h_pt = y1 - y0
            fs = max(5.0, round(h_pt * 0.90, 1))
            pad = 2.0
            sh = s.shapes.add_textbox(Emu(int((x0 - pad) * EMU_PER_PT)),
                                      Emu(int((y0 - pad * 0.6) * EMU_PER_PT)),
                                      Emu(int((x1 - x0 + pad * 2.4) * EMU_PER_PT)),
                                      Emu(int((h_pt + pad * 1.2) * EMU_PER_PT)))
            tf = sh.text_frame
            tf.word_wrap = False
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]; p.alignment = PP_ALIGN.LEFT
            r = p.add_run(); r.text = txt
            r.font.size = Pt(fs)
            r.font.color.rgb = RGBColor(*fg)
            r.font.bold = (ink > 0.30)
            _setfont(r)

    prs.save(out)
    return True


if __name__ == '__main__':
    import sys
    from coverage_benchmark import map_excel_to_report
    if len(sys.argv) < 2:
        print('사용법: python report_pptx.py <엑셀> [고객명]'); sys.exit(1)
    _cn = sys.argv[2] if len(sys.argv) > 2 else '고객'
    _rep = map_excel_to_report(sys.argv[1], settings={'client': _cn, 'branch': '온빛센터 바름지점',
                                                      'manager': '최은혜', 'title': '지점장', 'phone': ''})
    build_report_pptx(_rep, f'보장진단서_{_cn}.pptx')
    print('진단서 PPT 생성 완료(뼈대 유지 · 값 칸만 편집)')
